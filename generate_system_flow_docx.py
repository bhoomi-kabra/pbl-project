# Save this script as generate_system_flow_docx.py
# Prerequisite: pip install python-docx

import docx
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def build_system_flow_docx():
    doc = Document()

    # Set 1-inch margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Style Helper Functions
    def set_cell_background(cell, fill_hex):
        tcPr = cell._element.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        tcPr.append(shd)

    def add_title_block(text, size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6):
        p = doc.add_paragraph()
        p.alignment = align
        run = p.add_run(text)
        run.bold = bold
        run.font.name = 'Calibri'
        run.font.size = Pt(size)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.space_before = Pt(0)
        return p

    def add_h1(text):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = True
        run.font.name = 'Calibri'
        run.font.size = Pt(15)
        run.font.color.rgb = RGBColor(0, 51, 102) # Dark Navy
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)

    def add_h2(text):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = True
        run.font.name = 'Calibri'
        run.font.size = Pt(12.5)
        run.font.color.rgb = RGBColor(31, 78, 121)
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)

    def add_p(text, bold_prefix=None):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.bold = True
            r_pre.font.name = 'Calibri'
            r_pre.font.size = Pt(11)
        r = p.add_run(text)
        r.font.name = 'Calibri'
        r.font.size = Pt(11)

    def add_bullet(text, bold_prefix=None):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.bold = True
            r_pre.font.name = 'Calibri'
            r_pre.font.size = Pt(11)
        r = p.add_run(text)
        r.font.name = 'Calibri'
        r.font.size = Pt(11)

    def add_code_block(code_text):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        cell.width = Inches(6.5)
        set_cell_background(cell, "F4F6F9")
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(code_text)
        run.font.name = 'Consolas'
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(30, 30, 30)
        
        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    def add_styled_table(headers, data):
        table = doc.add_table(rows=len(data) + 1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = 'Table Grid'

        # Header Formatting
        hdr_cells = table.rows[0].cells
        for idx, text in enumerate(headers):
            hdr_cells[idx].text = text
            set_cell_background(hdr_cells[idx], "003366")
            p = hdr_cells[idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.bold = True
                run.font.name = 'Calibri'
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(10)

        # Row Data Formatting
        for r_idx, row_data in enumerate(data):
            row_cells = table.rows[r_idx + 1].cells
            bg_color = "F9FAFB" if r_idx % 2 == 1 else "FFFFFF"
            for c_idx, cell_value in enumerate(row_data):
                row_cells[c_idx].text = str(cell_value)
                set_cell_background(row_cells[c_idx], bg_color)
                p = row_cells[c_idx].paragraphs[0]
                for run in p.runs:
                    run.font.name = 'Calibri'
                    run.font.size = Pt(9.5)

        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # --- DOCUMENT CONTENT GENERATION ---

    # Institutional Header Block
    add_title_block("DEPARTMENT OF COMPUTER ENGINEERING", size=13)
    add_title_block("K. K. WAGH INSTITUTE OF ENGINEERING EDUCATION & RESEARCH, NASHIK", size=12)
    add_title_block("Academic Year: 2026–2027 | Project Based Learning (PBL) | Group No.: 7", size=10, bold=False, space_after=12)

    # Document Main Title
    add_title_block("TECHNICAL FLOW SPECIFICATION: NASHIK ROADS & CIVIC MONITOR", size=15)
    add_title_block("End-to-End Architectural Pipeline from Hazard Detection to SLA Resolution Audit", size=11, bold=False, space_after=18)

    # Section 1: Executive Architectural Overview
    add_h1("1. Executive System Flow Diagram")
    add_p("The ASCII architectural schematic below details the complete 7-step execution lifecycle of a civic grievance within the Nashik Roads & Civic Monitor ecosystem:")

    flow_ascii = """====================================================================================================
                       NASHIK ROADS & CIVIC MONITOR : MASTER SYSTEM FLOW
====================================================================================================

 [ STEP 1: HAZARD DETECTION & USER INPUT ]
   ├── Option A: Citizen types text/photo in AI Chatbot ("Sparking wire near K.K. Wagh, Panchavati")
   ├── Option B: Citizen selects location directly on Leaflet GIS Map (College Rd, Gangapur Rd)
   └── Option C: Weather Engine triggers Monsoon Risk Alert or Live Streamer simulates report
                                       │
                                       ▼
 [ STEP 2: AI INTENT CLASSIFICATION & DEEP-LINKING ]
   ├── AI Intent Model classifies hazard: ELECTRICAL_HAZARD (Confidence: 98%)
   ├── Auto-detects NMC Ward: Panchavati (Ward 1)
   └── Generates In-App Deep Link: 👉 Go to /app/complaints/electrical
                                       │
                                       ▼
 [ STEP 3: AUTO-FILLED COMPLAINT FORM & SUBMISSION ]
   ├── Pre-populates Category, Location Address, Ward Number, & Geotagged Evidence Photo
   └── Citizen clicks "Submit Complaint to NMC Grid"
                                       │
                                       ▼
 [ STEP 4: REST API & POSTGRESQL DATABASE PERSISTENCE ]
   ├── HTTP POST /api/complaints sent to Node.js Express API Server (Port 5000)
   ├── SQL INSERT query executed on Supabase Cloud PostgreSQL Database
   └── PostGIS extension stores spatial GPS coordinates [20.0180° N, 73.8180° E]
                                       │
                                       ▼
 [ STEP 5: REAL-TIME WEBSOCKETS BROADCAST (SOCKET.IO) ]
   ├── Server emits 'new_complaint' WebSocket event to all connected web browsers
   ├── Leaflet GIS Map drops an animated marker pin (🔴 Trenching / 🟡 Concreting / ⚡ Hazard)
   └── React UI updates live without page refresh (+1 KPI Counter & Live Toast Alert Banner)
                                       │
                                       ▼
 [ STEP 6: MUNICIPAL WORK EXECUTION & DLP TRACKING ]
   ├── Assigned Ward Engineer updates project lifecycle (TRENCHING → CONCRETING → CURING)
   ├── System tracks Contractor Defect Liability Period (DLP) (e.g. L&T Infra - 36 Months DLP)
   └── Ward Engineer uploads geotagged "After Repair" evidence photo
                                       │
                                       ▼
 [ STEP 7: "CLOSED != RESOLVED" CITIZEN VERIFICATION AUDIT ]
   ├── Ticket status set to EVIDENCE_UPLOADED
   ├── Side-by-Side Geotagged Photo Audit rendered: "Before Repair" vs "After Repair"
   └── Citizen casts audit vote:
         │
         ├── 🟢 "I CONFIRM IT'S FIXED"
         │     ├── Status updated to CLOSED_VERIFIED in Supabase PostgreSQL
         │     ├── SLA Resolution Rate updated on KPI Banner
         │     └── Citizen awarded +50 Civic Score Points
         │
         └── 🔴 "NOT FIXED PROPERLY - REOPEN & ESCALATE"
               ├── Status updated to REOPENED_ESCALATED
               ├── Triggers high-priority SLA Escalation Alert to NMC Chief Engineer
               └── Logs penalty warning against Contractor DLP Security Deposit
===================================================================================================="""
    
    add_code_block(flow_ascii)

    # Section 2: Detailed Technical Breakdowns
    add_h1("2. Technical Breakdown of System Steps")

    add_h2("Step 1: Hazard Detection & Multi-Channel User Input")
    add_bullet("Conversational NLP processing of natural text inputs and camera photo submissions via the floating chatbot widget.", bold_prefix="Option A (AI Chatbot Intake): ")
    add_bullet("Interactive Leaflet.js click-to-pin geolocation selection across Nashik's 6 administrative wards.", bold_prefix="Option B (GIS Direct Selection): ")
    add_bullet("Automated triggers driven by live OpenWeatherMap precipitation feeds or automated test script generators.", bold_prefix="Option C (System Alerts/Simulators): ")

    add_h2("Step 2: AI Intent Classification, Geofencing & Deep-Routing")
    add_bullet("Analyzes input keywords and photo features to output domain category (e.g., ELECTRICAL_HAZARD) with confidence scoring (>95%).", bold_prefix="Intent Model Execution: ")
    add_bullet("Executes PostGIS point-in-polygon queries against stored Nashik ward GeoJSON boundaries to assign Ward ID automatically.", bold_prefix="Automated Ward Binding: ")
    add_bullet("Generates direct client-side sub-route navigation (/app/complaints/electrical) without navigating external municipal websites.", bold_prefix="In-App Deep Routing: ")

    add_h2("Step 3: Auto-Filled Form Generation & Client Validation")
    add_p("The deep-link route automatically initializes the intake form, pre-filling verified parameters including Category, Extracted Address, Ward Number, and embedded EXIF GPS metadata. Client-side validation ensures photo payloads are compressed before network transfer.")

    add_h2("Step 4: Backend REST API & Spatial Database Persistence")
    add_bullet("Express API endpoint receiving JSON payloads containing text descriptions, base64/S3 photo links, and GPS floating-point pairs.", bold_prefix="HTTP POST /api/complaints: ")
    add_bullet("Relational schema persistence writing ticket records into Supabase PostgreSQL.", bold_prefix="Database Execution: ")
    add_bullet("Executes ST_SetSRID(ST_MakePoint(lng, lat), 4326) to store native spatial points for GIS proximity calculations.", bold_prefix="PostGIS Spatial Indexing: ")

    add_h2("Step 5: Real-Time WebSocket Broadcasting Engine")
    add_p("Upon successful database insertion, Socket.io broadcasts a 'new_complaint' payload to connected clients across the network. The Leaflet GIS map dynamically drops color-coded animated pins while KPI counters increment instantly without requiring page reloads.")

    add_h2("Step 6: Field Execution & Contractor DLP Tracking")
    add_p("Assigned Ward Engineers manage lifecycle stages via their mobile field console (TRENCHING → CONCRETING → CURING). The system tracks the Defect Liability Period (DLP) for the contractor agency and requires a compulsory geotagged 'After Repair' photo upload before advancing the ticket.")

    add_h2("Step 7: Citizen Verification Gate ('Closed != Resolved')")
    add_p("The system moves the ticket state to EVIDENCE_UPLOADED and presents a side-by-side audit view to the reporting citizen:")
    add_bullet("Closes ticket as CLOSED_VERIFIED, updates KPI SLA resolution metrics, and allocates +50 civic gamification score points to the user.", bold_prefix="Action: 'I Confirm It's Fixed': ")
    add_bullet("Transitions ticket to REOPENED_ESCALATED, generates a high-priority alert to the NMC Chief Engineer, and records a non-compliance penalty flag against the contractor's security deposit.", bold_prefix="Action: 'Not Fixed Properly': ")

    # Section 3: Data Flow & Endpoint Matrix Table
    add_h1("3. System Data Flow Matrix")
    add_styled_table(
        ["Step", "Endpoint / Handler", "Data Payload / Entity", "Target System / Store"],
        [
            ["Step 1-2", "POST /api/chatbot/parse", "{ text, photo_base64 }", "AI Intent Engine / NLP Module"],
            ["Step 3-4", "POST /api/complaints", "{ category, ward_id, lat, lng, image_url }", "Express API -> PostgreSQL (PostGIS)"],
            ["Step 5", "WebSocket Emit ('new_complaint')", "{ ticket_id, geometry, status, category }", "Socket.io Server -> React Client UI"],
            ["Step 6", "PATCH /api/projects/:id/status", "{ stage, engineer_id, proof_photo_url }", "Ward Engineer Console -> DB Store"],
            ["Step 7", "POST /api/complaints/:id/audit", "{ vote: 'CONFIRM' | 'REOPEN', comments }", "Citizen Verification Gate -> Escalation Engine"]
        ]
    )

    # Save to disk
    output_filename = "Nashik_Roads_Master_System_Flow.docx"
    doc.save(output_filename)
    print(f"Document successfully created: {output_filename}")

if __name__ == "__main__":
    build_system_flow_docx()